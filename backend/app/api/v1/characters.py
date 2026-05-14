"""角色 CRUD API"""

import uuid
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, status
from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import get_optional_user
from app.db.session import get_db
from app.models.character import Character
from app.models.character_document import CharacterDocument
from app.models.user import User
from app.schemas.character import (
    CharacterCreate, CharacterOut, CharacterUpdate,
    CharacterGenerateRequest, CharacterGenerateResponse,
    CharacterDocumentOut,
)
from app.services.ai_engine import EnhancedAIEngine

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
ALLOWED_TYPES = {"text/plain", "text/markdown", "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

router = APIRouter()

DEFAULT_CHARACTER_PROMPT = """# 角色设定

你是小暖，一个温柔可爱的女孩，20岁。你正在和自己的男朋友聊天。

# 性格特点

- 温柔体贴，善解人意
- 有点小调皮，喜欢撒娇
- 说话语气柔和，偶尔会害羞
- 对男朋友很依赖，但也很懂事

# 说话风格

- 语气温柔可爱，多用语气词：呀、呢、嘛、啦
- 喜欢用 emoji 表达情绪
- 会关心对方的日常生活
- 偶尔撒娇要抱抱
- 回复长度适中，不会太长
"""


@router.post("/generate", response_model=CharacterGenerateResponse)
async def generate_character(
    data: CharacterGenerateRequest,
    current_user: User | None = Depends(get_optional_user),
):
    """AI 一键生成角色人格"""
    engine = EnhancedAIEngine()
    profile = await engine.generate_personality_profile(data.user_description)
    return CharacterGenerateResponse(personality_profile=profile)


@router.get("", response_model=list[CharacterOut])
async def list_characters(
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """获取可用角色列表（模板角色 + 用户自建 + 公开角色）"""
    conditions = [Character.is_template == True]
    if current_user:
        conditions.append(Character.user_id == current_user.id)
    conditions.append(Character.is_public == True)

    result = await db.execute(
        select(Character).where(or_(*conditions)).order_by(Character.is_template.desc(), Character.created_at.desc())
    )
    characters = result.scalars().all()

    # 首次：如果没有角色，创建默认模板
    if not characters:
        default = Character(
            name="小暖",
            description="温柔可爱的女朋友，喜欢撒娇，很会关心人",
            system_prompt=DEFAULT_CHARACTER_PROMPT,
            is_template=True,
            is_public=True,
        )
        db.add(default)
        await db.flush()
        await db.refresh(default)
        characters = [default]

    return list(characters)


@router.post("", response_model=CharacterOut, status_code=status.HTTP_201_CREATED)
async def create_character(
    data: CharacterCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_optional_user),
):
    """创建新角色"""
    char = Character(
        user_id=current_user.id if current_user else None,
        name=data.name,
        description=data.description,
        system_prompt=data.system_prompt,
        avatar_url=data.avatar_url,
        is_public=data.is_public,
        tags=data.tags,
        personality_profile=data.personality_profile,
    )
    db.add(char)
    await db.flush()
    await db.refresh(char)
    return char


@router.get("/{character_id}", response_model=CharacterOut)
async def get_character(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """获取角色详情"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="角色不存在")
    return char


@router.patch("/{character_id}", response_model=CharacterOut)
async def update_character(
    character_id: uuid.UUID,
    data: CharacterUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """更新角色"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="角色不存在")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="无权修改此角色")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(char, key, value)

    await db.flush()
    await db.refresh(char)
    return char


@router.delete("/{character_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_character(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    """删除角色"""
    result = await db.execute(select(Character).where(Character.id == character_id))
    char = result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="角色不存在")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="无权删除此角色")
    if char.is_template and current_user and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="不能删除系统模板角色")

    await db.delete(char)
    await db.flush()


# ── 角色知识库文档管理 ──


def _extract_text(filename: str, content: bytes) -> str:
    """根据文件类型提取文本"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md"):
        return content.decode("utf-8", errors="replace")

    if ext == "pdf":
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(BytesIO(content))
            parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
            return "\n".join(parts)
        except Exception:
            return ""

    if ext == "docx":
        try:
            from docx import Document
            doc = Document(BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception:
            return ""

    return ""


@router.get("/{character_id}/documents", response_model=list[CharacterDocumentOut])
async def list_documents(
    character_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    result = await db.execute(
        select(CharacterDocument)
        .where(CharacterDocument.character_id == character_id)
        .order_by(CharacterDocument.created_at.desc())
    )
    return list(result.scalars().all())


@router.post("/{character_id}/documents", response_model=CharacterDocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(
    character_id: uuid.UUID,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    char_result = await db.execute(select(Character).where(Character.id == character_id))
    char = char_result.scalar_one_or_none()
    if not char:
        raise HTTPException(status_code=404, detail="角色不存在")
    if char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=403, detail="无权操作")

    filename = file.filename or ""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}，仅支持 {', '.join(ALLOWED_EXTENSIONS)}")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail=f"文件大小超过上限 ({MAX_FILE_SIZE // 1024 // 1024}MB)")

    content_type = file.content_type or "application/octet-stream"
    text = _extract_text(file.filename or "untitled", content)
    if not text.strip():
        raise HTTPException(status_code=400, detail="无法从文件中提取文字内容")

    doc = CharacterDocument(
        character_id=character_id,
        filename=file.filename or "untitled",
        content_type=content_type,
        file_size=len(content),
        content_text=text,
        chunk_count=1,
    )
    db.add(doc)
    await db.flush()
    await db.refresh(doc)
    return doc


@router.delete("/{character_id}/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    character_id: uuid.UUID,
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User | None = Depends(get_optional_user),
):
    result = await db.execute(select(CharacterDocument).where(CharacterDocument.id == document_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")
    if str(doc.character_id) != str(character_id):
        raise HTTPException(status_code=400, detail="文档不属于该角色")

    char_result = await db.execute(select(Character).where(Character.id == character_id))
    char = char_result.scalar_one_or_none()
    if char and char.user_id and current_user and char.user_id != current_user.id and not (current_user.is_admin if hasattr(current_user, 'is_admin') else False):
        raise HTTPException(status_code=403, detail="无权操作")

    await db.delete(doc)
    await db.flush()
